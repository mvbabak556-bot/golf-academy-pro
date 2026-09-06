# -*- coding: utf-8 -*-
"""موتور ساخت پرزنتیشن‌های فارسی گلف - راست‌چین (RTL) با پشتیبانی کامل."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

BASE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(BASE, "images")
CROP = os.path.join(IMG, "crop")
os.makedirs(CROP, exist_ok=True)

EMU_W, EMU_H = Inches(13.333), Inches(7.5)
SW, SH = 13.333, 7.5
FONT = "Vazirmatn"

# ---- academy branding ----
BRAND = "Puttclub"
BRAND_FA = "آکادمی گلف پات‌کلاب"
INSTRUCTOR = "الهام طالبی"
INSTRUCTOR_ROLE = "مدرس درجه یک گلف ایران"

# ---------- color helpers ----------
def C(hexs):
    hexs = hexs.lstrip("#")
    return RGBColor(int(hexs[0:2], 16), int(hexs[2:4], 16), int(hexs[4:6], 16))

WHITE = C("FFFFFF")
INK = C("1F2A24")
DARK = C("16211C")
MUTE = C("5C6B62")
CREAM = C("F6F8F3")
LINE = C("DCE5DA")

# ---------- image cover-crop ----------
def cover_crop(src, aspect, vbias=0.5):
    path = src if os.path.isabs(src) else os.path.join(IMG, src)
    im = Image.open(path).convert("RGB")
    w, h = im.size
    ia = w / h
    if abs(ia - aspect) < 0.01:
        out = os.path.join(CROP, os.path.basename(path))
        im.save(out, quality=90)
        return out
    if ia > aspect:  # too wide -> crop sides
        new_w = int(h * aspect)
        left = int((w - new_w) * vbias)
        box = (left, 0, left + new_w, h)
    else:  # too tall -> crop top/bottom
        new_h = int(w / aspect)
        top = int((h - new_h) * vbias)
        box = (0, top, w, top + new_h)
    im2 = im.crop(box)
    out = os.path.join(CROP, os.path.splitext(os.path.basename(path))[0] + f"_{aspect:.2f}.jpg")
    im2.save(out, quality=90)
    return out

# ---------- low-level ----------
def _set_font(run, size, bold, color, name=FONT, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:cs", "a:latin"):
        for el in rPr.findall(qn(tag)):
            rPr.remove(el)
    rPr.append(rPr.makeelement(qn("a:latin"), {"typeface": name}))
    rPr.append(rPr.makeelement(qn("a:cs"), {"typeface": name}))

def _rtl(p):
    pPr = p._p.get_or_add_pPr()
    pPr.set("rtl", "1")

def no_shadow(shape):
    shape.shadow.inherit = False

def add_rect(slide, x, y, w, h, fill, line=None, round_=False, shadow=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if round_:
        try:
            shp.adjustments[0] = 0.06
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    if not shadow:
        no_shadow(shp)
    return shp

def add_alpha(shape, color, opacity):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    srgb = shape.fill.fore_color._xFill.find(qn("a:srgbClr"))
    srgb.append(srgb.makeelement(qn("a:alpha"), {"val": str(int(opacity * 1000))}))
    no_shadow(shape)
    shape.line.fill.background()

def add_pic(slide, src, x, y, w, h, vbias=0.5, line=None):
    aspect = w / h
    p = cover_crop(src, aspect, vbias)
    pic = slide.shapes.add_picture(p, Inches(x), Inches(y), Inches(w), Inches(h))
    if line:
        pic.line.color.rgb = line
        pic.line.width = Pt(2)
    return pic

def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tb, tf

def para(tf, first=False, align=PP_ALIGN.RIGHT, space_after=8, space_before=0, line=1.08):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    _rtl(p)
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    try:
        p.line_spacing = line
    except Exception:
        pass
    return p

def run(p, text, size, bold, color, italic=False):
    r = p.add_run()
    r.text = text
    _set_font(r, size, bold, color, italic=italic)
    return r

# ---------- bullet list ----------
def add_bullets(slide, x, y, w, items, size=18, color=INK, gap=10, accent=None,
                lead_color=None, line=1.12, bullet="●", bullet_gap="  "):
    tb, tf = textbox(slide, x, y, w, 0.0)
    first = True
    for it in items:
        lvl = 0
        txt = it
        if isinstance(it, tuple):
            txt, lvl = it[0], it[1]
        p = para(tf, first=first, align=PP_ALIGN.RIGHT, space_after=gap, line=line)
        first = False
        if lvl >= 1:
            rb = run(p, "–" + bullet_gap, size - 4, True, accent or MUTE)
            run(p, txt, size - 3, False, color)
            p.level = 1
        else:
            run(p, bullet + bullet_gap, size - 5, True, accent or MUTE)
            # bold lead if "lead: rest"
            if "::" in txt:
                lead, rest = txt.split("::", 1)
                run(p, lead.strip(), size, True, lead_color or color)
                if rest.strip():
                    run(p, "  —  " + rest.strip(), size, False, color)
            else:
                run(p, txt, size, False, color)
    return tb

# ---------- scaffold ----------
class Deck:
    def __init__(self, accent, accent2, name_fa, name_en, level_fa):
        self.prs = Presentation()
        self.prs.slide_width = EMU_W
        self.prs.slide_height = EMU_H
        self.blank = self.prs.slide_layouts[6]
        self.accent = C(accent)
        self.accent2 = C(accent2)
        self.name_fa = name_fa
        self.name_en = name_en
        self.level_fa = level_fa
        self.page = 0

    def new_slide(self, bg=CREAM):
        s = self.prs.slides.add_slide(self.blank)
        add_rect(s, -0.1, -0.1, SW + 0.2, SH + 0.2, bg)
        self.page += 1
        return s

    def footer(self, s, dark_bg=False):
        col = WHITE if dark_bg else MUTE
        tb, tf = textbox(s, 0.55, SH - 0.5, 7.5, 0.35, anchor=MSO_ANCHOR.MIDDLE)
        p = para(tf, first=True, align=PP_ALIGN.LEFT, space_after=0)
        run(p, f"⛳ {BRAND}", 11, True, self.accent if not dark_bg else self.accent2)
        run(p, f"   ·   {self.name_en}", 10, False, col)
        tb2, tf2 = textbox(s, SW - 2.2, SH - 0.5, 1.6, 0.35, anchor=MSO_ANCHOR.MIDDLE)
        p2 = para(tf2, first=True, align=PP_ALIGN.LEFT, space_after=0)
        run(p2, str(self.page), 11, True, self.accent if not dark_bg else self.accent2)

    # ---- cover ----
    def cover(self, img, kicker, title, subtitle, meta_lines):
        s = self.new_slide(WHITE)
        add_pic(s, img, 0, 0, SW, SH, vbias=0.45)
        ov = add_rect(s, 0, 0, SW, SH, DARK)
        add_alpha(ov, DARK, 0.62)
        # bottom gradient band
        band = add_rect(s, 0, SH - 3.35, SW, 3.35, DARK)
        add_alpha(band, DARK, 0.55)
        # accent tab
        add_rect(s, SW - 3.2, 0.0, 3.2, 0.16, self.accent)
        # brand lockup (top-left)
        tb, tf = textbox(s, 0.7, 0.5, 5.0, 0.95, anchor=MSO_ANCHOR.TOP)
        p = para(tf, first=True, align=PP_ALIGN.LEFT, space_after=0)
        run(p, BRAND, 30, True, WHITE)
        p = para(tf, align=PP_ALIGN.LEFT, space_after=0)
        run(p, BRAND_FA + "  ·  " + INSTRUCTOR, 12.5, False, C("D7E2D6"))
        # instructor pill (top-right)
        pill = add_rect(s, SW - 4.55, 0.55, 3.85, 0.92, WHITE, round_=True)
        add_alpha(pill, WHITE, 0.12)
        tb, tf = textbox(s, SW - 4.55, 0.6, 3.85, 0.82, anchor=MSO_ANCHOR.MIDDLE)
        p = para(tf, first=True, align=PP_ALIGN.CENTER, space_after=1)
        run(p, "مدرس: " + INSTRUCTOR, 14.5, True, WHITE)
        p = para(tf, align=PP_ALIGN.CENTER, space_after=0)
        run(p, INSTRUCTOR_ROLE, 11, False, self.accent2)
        tb, tf = textbox(s, 1.0, SH - 3.05, SW - 2.0, 2.6, anchor=MSO_ANCHOR.MIDDLE)
        p = para(tf, first=True, align=PP_ALIGN.RIGHT, space_after=6)
        run(p, kicker, 17, True, self.accent2)
        p = para(tf, align=PP_ALIGN.RIGHT, space_after=6)
        run(p, title, 46, True, WHITE)
        p = para(tf, align=PP_ALIGN.RIGHT, space_after=14)
        run(p, subtitle, 20, False, C("D7E2D6"))
        for i, ml in enumerate(meta_lines):
            p = para(tf, align=PP_ALIGN.RIGHT, space_after=3)
            run(p, "●  ", 11, True, self.accent2)
            run(p, ml, 14.5, False, C("C9D6C8"))
        self.footer(s, dark_bg=True)
        return s

    # ---- section / objectives hero ----
    def section(self, no, title, subtitle=None, items=None):
        s = self.new_slide(self.accent)
        add_rect(s, 0, 0, 0.35, SH, self.accent2)
        # big faint number
        tb, tf = textbox(s, 0.6, 0.5, 4.5, 3.2, anchor=MSO_ANCHOR.TOP)
        p = para(tf, first=True, align=PP_ALIGN.LEFT, space_after=0)
        run(p, no, 150, True, WHITE)
        _set_alpha_run(p)
        tb, tf = textbox(s, 1.2, 2.05, SW - 2.4, 1.5, anchor=MSO_ANCHOR.MIDDLE)
        p = para(tf, first=True, align=PP_ALIGN.RIGHT, space_after=8)
        run(p, title, 40, True, WHITE)
        if subtitle:
            p = para(tf, align=PP_ALIGN.RIGHT, space_after=0)
            run(p, subtitle, 19, False, C("EAF3E8"))
        if items:
            add_bullets(s, SW - 6.4, 3.75, 5.4, items, size=16,
                        color=WHITE, lead_color=WHITE, accent=self.accent2, gap=9)
        self.footer(s, dark_bg=True)
        return s

    # ---- header for content ----
    def header(self, s, kicker, title, img_side=False):
        # accent bar on right (RTL start)
        add_rect(s, SW - 0.28, 0, 0.28, SH, self.accent)
        tx = 0.7 if img_side else 0.9
        tw = (SW - 5.6) if img_side else (SW - 1.8)
        # title sits on the right side
        tb, tf = textbox(s, SW - 0.28 - tw - 0.55, 0.55, tw, 1.35, anchor=MSO_ANCHOR.TOP)
        p = para(tf, first=True, align=PP_ALIGN.RIGHT, space_after=4)
        run(p, kicker, 14, True, self.accent)
        p = para(tf, align=PP_ALIGN.RIGHT, space_after=0, line=1.0)
        run(p, title, 29, True, INK)
        # underline
        add_rect(s, SW - 0.28 - 2.4, 1.72, 2.0, 0.045, self.accent2)
        return s

    # ---- content slide with left image panel ----
    def content_image(self, kicker, title, items, img, caption=None, size=17.5,
                      vbias=0.5, note=None):
        s = self.new_slide(WHITE)
        panel_w = 4.55
        add_pic(s, img, 0, 0, panel_w, SH, vbias=vbias)
        cap = add_rect(s, 0, SH - 0.95, panel_w, 0.95, DARK)
        add_alpha(cap, DARK, 0.45)
        if caption:
            tb, tf = textbox(s, 0.35, SH - 0.95, panel_w - 0.6, 0.95, anchor=MSO_ANCHOR.MIDDLE)
            p = para(tf, first=True, align=PP_ALIGN.LEFT, space_after=0)
            run(p, caption, 13, False, WHITE)
        self.header(s, kicker, title, img_side=True)
        tx = panel_w + 0.55
        tw = SW - tx - 0.65
        ay = 2.1
        add_bullets(s, tx, ay, tw, items, size=size, accent=self.accent,
                    lead_color=self.accent, gap=11)
        if note:
            ny = SH - 1.55
            card = add_rect(s, tx, ny, tw, 0.92, CREAM, round_=True)
            add_rect(s, tx, ny, 0.09, 0.92, self.accent2)
            tb, tf = textbox(s, tx + 0.3, ny + 0.1, tw - 0.55, 0.72, anchor=MSO_ANCHOR.MIDDLE)
            p = para(tf, first=True, align=PP_ALIGN.RIGHT, space_after=0, line=1.1)
            run(p, "نکته مربی:  ", 14, True, self.accent)
            run(p, note, 14, False, INK)
        self.footer(s)
        return s

    # ---- content slide no image ----
    def content(self, kicker, title, items, size=19, note=None, two_col=False):
        s = self.new_slide(WHITE)
        self.header(s, kicker, title)
        tx = 0.9
        tw = SW - 1.8
        ay = 2.15
        if two_col and len(items) > 6:
            half = (len(items) + 1) // 2
            colw = (tw - 0.6) / 2
            add_bullets(s, tx, ay, colw, items[:half], size=size - 1, accent=self.accent,
                        lead_color=self.accent, gap=11)
            add_bullets(s, tx + colw + 0.6, ay, colw, items[half:], size=size - 1,
                        accent=self.accent, lead_color=self.accent, gap=11)
        else:
            add_bullets(s, tx, ay, tw, items, size=size, accent=self.accent,
                        lead_color=self.accent, gap=12)
        if note:
            ny = SH - 1.5
            card = add_rect(s, tx, ny, tw, 0.9, CREAM, round_=True)
            add_rect(s, tx, ny, 0.09, 0.9, self.accent2)
            tb, tf = textbox(s, tx + 0.3, ny + 0.1, tw - 0.55, 0.7, anchor=MSO_ANCHOR.MIDDLE)
            p = para(tf, first=True, align=PP_ALIGN.RIGHT, space_after=0, line=1.1)
            run(p, "💡 ", 14, True, self.accent)
            run(p, note, 14.5, True, INK)
        self.footer(s)
        return s

    # ---- cards (chips for clubs / terms) ----
    def cards(self, kicker, title, cards, cols=3, note=None):
        s = self.new_slide(WHITE)
        self.header(s, kicker, title)
        tx, ty = 0.85, 2.15
        tw = SW - 1.7
        rows = (len(cards) + cols - 1) // cols
        gap = 0.3
        cw = (tw - gap * (cols - 1)) / cols
        reserve = 1.55 if note else 1.0
        ch = min(1.55, (SH - ty - reserve - (rows - 1) * gap) / rows)
        if ch < 0.95:
            ch = max(0.9, (SH - ty - reserve - (rows - 1) * gap) / rows)
        for i, (head, body) in enumerate(cards):
            r, c = divmod(i, cols)
            # RTL: first card on right
            x = tx + (cols - 1 - c) * (cw + gap)
            y = ty + r * (ch + gap)
            card = add_rect(s, x, y, cw, ch, CREAM, round_=True)
            add_rect(s, x, y, cw, 0.1, self.accent)
            tb, tf = textbox(s, x + 0.22, y + 0.18, cw - 0.44, ch - 0.3, anchor=MSO_ANCHOR.MIDDLE)
            p = para(tf, first=True, align=PP_ALIGN.RIGHT, space_after=4, line=1.05)
            run(p, head, 16.5, True, self.accent)
            if body:
                p = para(tf, align=PP_ALIGN.RIGHT, space_after=0, line=1.08)
                run(p, body, 12.5, False, MUTE)
        if note:
            ny = SH - 1.2
            card = add_rect(s, tx, ny, tw, 0.72, self.accent, round_=True)
            tb, tf = textbox(s, tx + 0.3, ny + 0.08, tw - 0.6, 0.56, anchor=MSO_ANCHOR.MIDDLE)
            p = para(tf, first=True, align=PP_ALIGN.RIGHT, space_after=0)
            run(p, note, 14.5, True, WHITE)
        self.footer(s)
        return s

    # ---- checklist ----
    def checklist(self, kicker, title, groups, note=None):
        s = self.new_slide(WHITE)
        self.header(s, kicker, title)
        tx, ty = 0.85, 2.1
        tw = SW - 1.7
        n = len(groups)
        gap = 0.4
        cw = (tw - gap * (n - 1)) / n
        ch = SH - ty - 1.15
        for i, (head, items) in enumerate(groups):
            x = tx + (n - 1 - i) * (cw + gap)
            card = add_rect(s, x, ty, cw, ch, CREAM, round_=True)
            add_rect(s, x, ty, cw, 0.72, self.accent, round_=False)
            add_rect(s, x, ty + 0.62, cw, 0.1, self.accent2)
            tb, tf = textbox(s, x + 0.2, ty + 0.06, cw - 0.4, 0.62, anchor=MSO_ANCHOR.MIDDLE)
            p = para(tf, first=True, align=PP_ALIGN.RIGHT, space_after=0)
            run(p, head, 16, True, WHITE)
            tb, tf = textbox(s, x + 0.28, ty + 0.95, cw - 0.5, ch - 1.1, anchor=MSO_ANCHOR.TOP)
            first = True
            for it in items:
                p = para(tf, first=first, align=PP_ALIGN.RIGHT, space_after=9, line=1.08)
                first = False
                run(p, "☐  ", 13, True, self.accent)
                run(p, it, 13.5, False, INK)
        if note:
            ny = SH - 0.85
            tb, tf = textbox(s, tx, ny, tw, 0.5, anchor=MSO_ANCHOR.MIDDLE)
            p = para(tf, first=True, align=PP_ALIGN.RIGHT, space_after=0)
            run(p, note, 14, True, self.accent)
        self.footer(s)
        return s

    # ---- flow / path ----
    def flow(self, kicker, title, steps, note=None):
        s = self.new_slide(WHITE)
        self.header(s, kicker, title)
        tx, ty = 0.9, 2.9
        tw = SW - 1.8
        n = len(steps)
        gap = 0.25
        bw = (tw - gap * (n - 1)) / n
        bh = 1.5
        for i, st in enumerate(steps):
            # RTL: step 1 on the right
            x = tx + (n - 1 - i) * (bw + gap)
            box = add_rect(s, x, ty, bw, bh, CREAM, round_=True)
            add_rect(s, x, ty, bw, 0.1, self.accent)
            # number circle
            circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + bw/2 - 0.32), Inches(ty - 0.42),
                                      Inches(0.64), Inches(0.64))
            circ.fill.solid(); circ.fill.fore_color.rgb = self.accent
            circ.line.fill.background(); no_shadow(circ)
            ctf = circ.text_frame; ctf.word_wrap = True
            cp = ctf.paragraphs[0]; cp.alignment = PP_ALIGN.CENTER
            cr = cp.add_run(); cr.text = str(i + 1); _set_font(cr, 18, True, WHITE)
            tb, tf = textbox(s, x + 0.12, ty + 0.35, bw - 0.24, bh - 0.45, anchor=MSO_ANCHOR.MIDDLE)
            label = st if isinstance(st, str) else st[0]
            sub = None if isinstance(st, str) else (st[1] if len(st) > 1 else None)
            p = para(tf, first=True, align=PP_ALIGN.CENTER, space_after=2, line=1.0)
            run(p, label, 14, True, INK)
            if sub:
                p = para(tf, align=PP_ALIGN.CENTER, space_after=0, line=1.0)
                run(p, sub, 10.5, False, MUTE)
        if note:
            ny = SH - 1.55
            card = add_rect(s, tx, ny, tw, 0.85, CREAM, round_=True)
            add_rect(s, tx, ny, 0.09, 0.85, self.accent2)
            tb, tf = textbox(s, tx + 0.3, ny + 0.1, tw - 0.55, 0.65, anchor=MSO_ANCHOR.MIDDLE)
            p = para(tf, first=True, align=PP_ALIGN.RIGHT, space_after=0, line=1.1)
            run(p, note, 15, True, INK)
        self.footer(s)
        return s

    # ---- table ----
    def table(self, kicker, title, headers, rows, note=None, col_align=None, fs=13):
        s = self.new_slide(WHITE)
        self.header(s, kicker, title)
        tx, ty = 0.9, 2.15
        tw = SW - 1.8
        nrows = len(rows) + 1
        ncols = len(headers)
        th = min(0.62, (SH - ty - (1.1 if note else 0.7)) / nrows)
        gtbl = s.shapes.add_table(nrows, ncols, Inches(tx), Inches(ty), Inches(tw), Inches(th * nrows)).table
        # turn off banding default style by setting fills manually
        try:
            gtbl.first_row = False
        except Exception:
            pass
        # column widths
        for c in range(ncols):
            gtbl.columns[c].width = Inches(tw / ncols)
        for r in range(nrows):
            gtbl.rows[r].height = Inches(th)
        for c, htext in enumerate(headers):
            cell = gtbl.cell(0, c)
            cell.fill.solid(); cell.fill.fore_color.rgb = self.accent
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_top = Pt(2); cell.margin_bottom = Pt(2)
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _rtl(p)
            rr = p.add_run(); rr.text = htext; _set_font(rr, fs + 1, True, WHITE)
        for ri, row in enumerate(rows, start=1):
            for c, val in enumerate(row):
                cell = gtbl.cell(ri, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = CREAM if ri % 2 else WHITE
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.margin_top = Pt(1); cell.margin_bottom = Pt(1)
                tf = cell.text_frame; tf.word_wrap = True
                p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _rtl(p)
                bold = (c == 0)
                col = self.accent if c == 0 else INK
                rr = p.add_run(); rr.text = str(val); _set_font(rr, fs, bold, col)
        if note:
            ny = ty + th * nrows + 0.2
            tb, tf = textbox(s, tx, ny, tw, 0.6, anchor=MSO_ANCHOR.MIDDLE)
            p = para(tf, first=True, align=PP_ALIGN.RIGHT, space_after=0)
            run(p, "●  ", 11, True, self.accent)
            run(p, note, 13.5, False, MUTE)
        self.footer(s)
        return s

    # ---- quote / golden rule ----
    def quote(self, kicker, title, big, small=None):
        s = self.new_slide(DARK)
        add_rect(s, 0, 0, SW, 0.18, self.accent)
        add_rect(s, 0, SH - 0.18, SW, 0.18, self.accent2)
        tb, tf = textbox(s, 1.4, 1.7, SW - 2.8, 1.0, anchor=MSO_ANCHOR.MIDDLE)
        p = para(tf, first=True, align=PP_ALIGN.RIGHT, space_after=0)
        run(p, kicker + "  |  " + title, 16, True, self.accent2)
        tb, tf = textbox(s, 1.3, 2.7, SW - 2.6, 2.6, anchor=MSO_ANCHOR.MIDDLE)
        p = para(tf, first=True, align=PP_ALIGN.CENTER, space_after=0, line=1.25)
        run(p, "« " + big + " »", 27, True, WHITE)
        if small:
            p = para(tf, align=PP_ALIGN.CENTER, space_before=18, line=1.2)
            run(p, small, 16, False, C("C9D6C8"))
        self.footer(s, dark_bg=True)
        return s

    # ---- closing / roadmap ----
    def closing(self, title, levels, final_note, img=None):
        s = self.new_slide(WHITE)
        if img:
            add_pic(s, img, 0, 0, SW, SH, vbias=0.5)
            ov = add_rect(s, 0, 0, SW, SH, DARK); add_alpha(ov, DARK, 0.72)
        else:
            add_rect(s, -0.1, -0.1, SW + 0.2, SH + 0.2, DARK)
        tb, tf = textbox(s, 1.0, 0.7, SW - 2.0, 0.9, anchor=MSO_ANCHOR.MIDDLE)
        p = para(tf, first=True, align=PP_ALIGN.CENTER, space_after=0)
        run(p, title, 30, True, WHITE)
        # level pills
        n = len(levels)
        pw = (SW - 2.0 - 0.4 * (n - 1)) / n
        y = 2.5
        for i, (lab, desc, colh) in enumerate(levels):
            x = 1.0 + (n - 1 - i) * (pw + 0.4)
            card = add_rect(s, x, y, pw, 2.0, WHITE, round_=True)
            add_rect(s, x, y, pw, 0.62, C(colh), round_=False)
            tb, tf = textbox(s, x + 0.12, y + 0.05, pw - 0.24, 0.55, anchor=MSO_ANCHOR.MIDDLE)
            p = para(tf, first=True, align=PP_ALIGN.CENTER, space_after=0)
            run(p, lab, 16, True, WHITE)
            tb, tf = textbox(s, x + 0.18, y + 0.75, pw - 0.36, 1.1, anchor=MSO_ANCHOR.MIDDLE)
            p = para(tf, first=True, align=PP_ALIGN.CENTER, space_after=0, line=1.12)
            run(p, desc, 13.5, False, INK)
            # arrow
            if i < n - 1:
                ar = s.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, Inches(x - 0.42), Inches(y + 0.75),
                                       Inches(0.34), Inches(0.5))
                ar.fill.solid(); ar.fill.fore_color.rgb = self.accent2
                ar.line.fill.background(); no_shadow(ar)
        tb, tf = textbox(s, 1.2, 4.85, SW - 2.4, 1.35, anchor=MSO_ANCHOR.MIDDLE)
        p = para(tf, first=True, align=PP_ALIGN.CENTER, space_after=10, line=1.25)
        run(p, final_note, 19, True, C("EAF3E8"))
        p = para(tf, align=PP_ALIGN.CENTER, space_after=0)
        run(p, f"⛳ {BRAND}  —  {BRAND_FA}   |   مدرس: {INSTRUCTOR} ({INSTRUCTOR_ROLE})",
            14, True, self.accent2 if img else self.accent)
        self.footer(s, dark_bg=True)
        return s

    def save(self, path):
        self.prs.save(path)
        try:
            from embed_fonts import embed_fonts
            embed_fonts(path)
        except Exception as e:
            print("  (font embed skipped:", e, ")")


def _set_alpha_run(p):
    """make the big faint number soft white"""
    for r in p.runs:
        rPr = r._r.get_or_add_rPr()
        srgb = rPr.find(qn("a:solidFill"))
        if srgb is not None:
            rPr.remove(srgb)
        fill = rPr.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": "FFFFFF"})
        clr.append(clr.makeelement(qn("a:alpha"), {"val": "22000"}))
        fill.append(clr)
        rPr.append(fill)
